"""Bounded in-memory text extraction. Originals are never copied to the workspace."""
import csv
import io
from pathlib import Path
import zipfile
from app.core.data_guard import Part

MAX_BYTES = 8 * 1024 * 1024
MAX_TEXT = 120000
SUFFIXES = {'.txt','.md','.csv','.xlsx','.docx','.pptx'}


def extract(filename, content):
    suffix = Path(filename).suffix.lower()
    if suffix not in SUFFIXES:
        raise ValueError('지원 형식: TXT, MD, CSV, XLSX, DOCX, PPTX. 구형·암호화 파일은 변환 후 등록하세요.')
    if len(content)>MAX_BYTES: raise ValueError('파일은 8 MiB 이하만 지원합니다.')
    parts, size, visited = [], 0, 0
    def add(location, text, field=''):
        nonlocal size, visited
        visited += 1
        if visited>100000: raise ValueError('항목 처리 한도를 초과했습니다.')
        text = str(text) if text is not None else ''
        if len(text)>16000: raise ValueError('한 셀·문단은 16,000자 이하만 지원합니다.')
        size += len(text)
        if size>MAX_TEXT or len(parts)>=10000: raise ValueError('추출 한도를 초과했습니다. 필요한 시트·페이지만 나누어 등록하세요.')
        if text.strip(): parts.append(Part(location,text,str(field or '')))
    omissions = []
    try:
        if suffix in {'.txt','.md','.csv'}:
            text = content.decode('utf-8-sig')
            if '\x00' in text: raise ValueError()
            if suffix=='.csv':
                reader = csv.reader(io.StringIO(text)); headers = next(reader, [])
                for col, value in enumerate(headers,1): add(f'R1C{col}',value)
                for row, values in enumerate(reader,2):
                    for col, value in enumerate(values,1): add(f'R{row}C{col}',value,headers[col-1] if col<=len(headers) else '')
            else:
                for index,line in enumerate(text.splitlines(),1): add(f'L{index}',line)
        else:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries)>3000 or sum(e.file_size for e in entries)>32*1024*1024 or any(e.flag_bits&1 for e in entries):
                    raise ValueError('압축 해제 한도를 초과했거나 암호화된 파일입니다.')
                # Reject entity declarations before passing XML to third-party parsers.
                for entry in entries:
                    if entry.filename.endswith(('.xml','.rels')):
                        xml = archive.read(entry).upper()
                        if b'\x00' in xml or b'<!DOCTYPE' in xml or b'<!ENTITY' in xml: raise ValueError('지원하지 않는 XML 구조입니다.')
            omissions = ['이미지·도형·내장 객체·메타데이터는 전달하지 않음']
            if suffix=='.xlsx':
                from openpyxl import load_workbook
                book = load_workbook(io.BytesIO(content),read_only=True,data_only=False,keep_links=False)
                try:
                    for sheet_no,sheet in enumerate(book.worksheets,1):
                        if sheet.sheet_state!='visible':
                            omissions.append(f'시트 {sheet_no}: 숨김 시트 제외'); continue
                        if (sheet.max_row or 0)>10000 or (sheet.max_column or 0)>1000 or (sheet.max_row or 0)*(sheet.max_column or 0)>100000: raise ValueError('시트 크기 한도를 초과했습니다.')
                        headers = []
                        for row_no,row in enumerate(sheet.iter_rows(),1):
                            if row_no>10000: raise ValueError('행 한도를 초과했습니다.')
                            if row_no==1: headers = [cell.value for cell in row]
                            for col,cell in enumerate(row,1):
                                if cell.data_type=='f':
                                    if '수식 셀 제외' not in omissions: omissions.append('수식 셀 제외')
                                    continue
                                add(f'S{sheet_no}!R{row_no}C{col}',cell.value,headers[col-1] if row_no>1 and col<=len(headers) else '')
                finally: book.close()
                omissions.append('주석 제외')
            elif suffix=='.docx':
                from docx import Document
                doc = Document(io.BytesIO(content))
                for i,p in enumerate(doc.paragraphs,1): add(f'P{i}',p.text)
                for t,table in enumerate(doc.tables,1):
                    headers = [c.text for c in table.rows[0].cells] if table.rows else []
                    for r,row in enumerate(table.rows,1):
                        for c,cell in enumerate(row.cells,1): add(f'T{t}R{r}C{c}',cell.text,headers[c-1] if r>1 and c<=len(headers) else '')
                omissions.append('머리말·꼬리말·주석·중첩 표·변경 이력 제외')
            else:
                from pptx import Presentation
                deck = Presentation(io.BytesIO(content))
                for s,slide in enumerate(deck.slides,1):
                    for n,shape in enumerate(slide.shapes,1):
                        if shape.has_text_frame: add(f'S{s}T{n}',shape.text)
                        if shape.has_table:
                            table = shape.table; headers = [c.text for c in table.rows[0].cells]
                            for r,row in enumerate(table.rows,1):
                                for c,cell in enumerate(row.cells,1): add(f'S{s}T{n}R{r}C{c}',cell.text,headers[c-1] if r>1 else '')
                omissions.append('발표자 노트·그룹 도형·마스터 제외')
    except (ValueError,UnicodeError) as exc:
        # Do not reflect parser messages which may include source data.
        raise ValueError('파일을 읽을 수 없거나 처리 한도를 초과했습니다. 텍스트는 UTF-8을 사용하세요.') from None
    except Exception:
        raise ValueError('파일 추출에 실패했습니다. 원본은 저장하거나 전달하지 않았습니다.') from None
    if not parts: raise ValueError('전달 가능한 텍스트가 없습니다.')
    return parts, omissions
