import io
import json
import socket
import zipfile
import pytest
from fastapi.testclient import TestClient
from app.main import create_app
from app.models.schemas import AgentResult
from app.core.data_guard import guard, Part, GuardOptions, FieldRule
from app.core.material_extract import extract, MAX_BYTES


def test_local_detectors_and_audit_do_not_keep_secrets(monkeypatch):
    def forbidden(*args, **kwargs): raise AssertionError('Network is forbidden')
    monkeypatch.setattr(socket.socket, 'connect', forbidden)
    values=['AKIA'+'A'*16, 'ghp_'+'B'*36, 'example-password-987', 'dbuser:example-pass']
    parts=[Part('L1',values[0]),Part('L2',values[1]),Part('L3','password='+values[2]),Part('L4','postgres://'+values[3]+'@localhost/demo')]
    report=guard(parts,GuardOptions())
    serialized=json.dumps(report)
    assert all(value not in serialized for value in values)
    assert len(report['findings'])>=4
    assert report['status']=='ready'


@pytest.mark.parametrize('level',['strong','medium','low','allow'])
def test_strengths_and_korean_fields(level):
    parts=[Part('R2C1','홍길동','이름'),Part('R2C2','010-1234-2324','휴대폰'),
           Part('R2C3','dkwm12345','아이디'),Part('R2C4','900101-1234567'),Part('L1','password=dummy-pass')]
    report=guard(parts,GuardOptions(level=level))
    text='\n'.join(p['text'] for p in report['parts'])
    if level=='medium': assert all(value in text for value in ['홍*동','010-****-2324','dkwm*****'])
    if level=='strong': assert '홍길동' not in text and '2324' not in text
    if level in ('low','allow'): assert '홍길동' in text and '010-1234-2324' in text
    assert ('dummy-pass' in text)==(level=='allow')
    assert ('900101-1234567' in text)==(level=='allow')
    assert 'dummy-pass' not in json.dumps(report['findings'])


def test_custom_fields_labeled_text_and_whole_private_key_block():
    result=guard([Part('L1','담당자: 홍길동')],GuardOptions(level='medium',fields=[FieldRule(field='담당자',kind='name')]))
    assert result['parts'][0]['text']=='담당자: 홍*동'
    parts=[Part('L1','-----BEGIN OPENSSH PRIVATE KEY-----'),Part('L2','dummy payload')]
    for level in ('strong','medium','low'):
        report=guard(parts,GuardOptions(level=level))
        assert report['status']=='blocked' and report['parts']==[]
    assert guard(parts,GuardOptions(level='allow'))['parts'][1]['text']=='dummy payload'


@pytest.mark.parametrize('suffix',['.csv','.xlsx','.docx','.pptx'])
def test_office_tables_use_headers_without_original_file_names(suffix):
    stream=io.BytesIO()
    if suffix=='.csv': stream.write('이름,휴대폰\n홍길동,010-1234-2324'.encode())
    elif suffix=='.xlsx':
        from openpyxl import Workbook
        book=Workbook();sheet=book.active;sheet.append(['이름','휴대폰']);sheet.append(['홍길동','010-1234-2324']);sheet.append(['=1+1',None]);book.save(stream)
    elif suffix=='.docx':
        from docx import Document
        doc=Document();table=doc.add_table(rows=2,cols=2)
        for row,values in zip(table.rows,[['이름','휴대폰'],['홍길동','010-1234-2324']]):
            for cell,value in zip(row.cells,values): cell.text=value
        doc.save(stream)
    else:
        from pptx import Presentation
        from pptx.util import Inches
        deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6]);table=slide.shapes.add_table(2,2,0,0, Inches(5),Inches(2)).table
        for r,values in enumerate([['이름','휴대폰'],['홍길동','010-1234-2324']]):
            for c,value in enumerate(values): table.cell(r,c).text=value
        deck.save(stream)
    parts,omissions=extract('홍길동'+suffix,stream.getvalue())
    report=guard(parts,GuardOptions(level='medium'))
    text=json.dumps(report,ensure_ascii=False)
    assert '홍*동' in text and '010-****-2324' in text and '홍길동' not in text
    if suffix!='.csv': assert omissions


def test_extract_fails_closed_for_unsupported_malformed_and_entities():
    for name,content in [('a.pdf',b'abc'),('a.txt',b'\xff'),('a.docx',b'not zip'),('a.txt',b'x'*(MAX_BYTES+1))]:
        with pytest.raises(ValueError): extract(name,content)
    stream=io.BytesIO()
    with zipfile.ZipFile(stream,'w') as archive: archive.writestr('document.xml',b'<!DOCTYPE fake>')
    with pytest.raises(ValueError): extract('a.docx',stream.getvalue())


class Recorder:
    def __init__(self): self.contexts=[]
    async def run(self, context):
        self.contexts.append(context)
        return AgentResult(adapter='mock',output='done')


def test_upload_run_audit_restart_policy_change_and_project_scope(tmp_path):
    agent=Recorder();root=tmp_path/'workspaces'
    with TestClient(create_app(root,agent)) as client:
        for name in ('one','two'): client.post('/api/projects',json={'name':name})
        base='/api/projects/one'
        client.put(base+'/guard',json={'level':'medium'})
        response=client.post(base+'/materials',params={'filename':'홍길동.csv','test_data':True},content='이름,비밀번호\n홍길동,dummy-pass'.encode())
        assert response.status_code==201,response.text
        material=response.json();identity=material['id']
        assert material['test_data'] and 'dummy-pass' not in response.text and '홍길동' not in response.text
        assert client.get('/api/projects/two/materials/'+identity).status_code==404
        result=client.post(base+'/commands',json={'text':'요약','task':'검증','material_ids':[identity]})
        assert result.status_code==200,result.text
        assert '홍*동' in agent.contexts[-1].files[0].content
        assert 'dummy-pass' not in agent.contexts[-1].model_dump_json()
        run_id=result.json()['run_id'];audit=client.get(base+'/runs/'+run_id).json()['metadata']['guard']
        assert audit['state']=='dispatch_attempted' and audit['materials'][0]['id']==identity
        client.put(base+'/guard',json={'level':'strong'})
        assert client.get(base+'/materials').json()['materials'][0]['stale']
        assert client.post(base+'/commands',json={'text':'again','material_ids':[identity]}).status_code==400
        private=client.post(base+'/materials?filename=key.txt',content=b'-----BEGIN PRIVATE KEY-----\nfake').json()
        assert private['status']=='blocked'
        assert client.post(base+'/commands',json={'text':'again','material_ids':[private['id']]}).status_code==400
        client.put(base+'/guard',json={'level':'allow'})
        raw=client.post(base+'/materials?filename=test.txt',content=b'password=dummy-pass').json()
        assert raw['findings'][0]['action']=='raw'
        assert 'dummy-pass' in client.get(base+'/materials/'+raw['id']).json()['content']
        result=client.post(base+'/commands',json={'text':'allow','material_ids':[raw['id']]})
        assert result.status_code==200 and 'dummy-pass' in agent.contexts[-1].files[0].content
        assert client.post(base+'/materials?filename=x.txt',content=b'x'*(MAX_BYTES+1)).status_code==400
    with TestClient(create_app(root)) as client:
        assert client.get(base+'/guard').json()['level']=='allow'
        assert client.get(base+'/runs/'+run_id).json()['metadata']['guard']==audit
        client.delete(base+'/materials/'+identity)
        assert client.get(base+'/materials/'+identity).status_code==404


@pytest.mark.parametrize('client_name',['codex','claude'])
def test_both_native_adapters_receive_only_guarded_attachment(tmp_path,monkeypatch,client_name):
    from app.llm.cli import CliAdapter
    monkeypatch.setattr('app.core.orchestrator.executable',lambda *args:'unused')
    received=[]
    async def execute(self,prompt,*args):
        received.append(prompt);return AgentResult(adapter=client_name,output='done')
    monkeypatch.setattr(CliAdapter,'execute',execute)
    with TestClient(create_app(tmp_path/'workspaces')) as client:
        client.post('/api/projects',json={'name':'one'})
        base='/api/projects/one'
        identity=client.post(base+'/materials?filename=x.txt',content=b'password=dummy-pass').json()['id']
        response=client.post(base+'/commands',json={'text':'test','client':client_name,'material_ids':[identity]})
        assert response.status_code==200,response.text
        assert 'dummy-pass' not in received[0] and '[비밀값]' in received[0]
