"""Selected Excel cells must be guarded before persistence or adapter dispatch."""
import hashlib
import io
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pydantic import ValidationError

from app.core.data_guard import GuardOptions, RULE_VERSION
from app.core.material_extract import extract
from app.core.materials import fingerprint
from app.main import create_app
from app.models.schemas import AgentResult


def workbook():
    book = Workbook()
    sheet = book.active
    sheet.title = 'private-sheet-title'
    sheet.append(['이름', 'internal-only-header', '휴대폰', '비밀번호'])
    sheet.append(['홍길동', 'excluded-column-marker', '010-1234-2324', 'dummy-pass'])
    other = book.create_sheet('other-private-title')
    other.append(['unselected-sheet-marker'])
    hidden = book.create_sheet('hidden-private-title')
    hidden.append(['hidden-cell-marker'])
    hidden.sheet_state = 'hidden'
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def test_selected_coordinates_and_headers_preserve_original_positions():
    options = GuardOptions(excel_sheets=[{'sheet': 1, 'columns': [3, 1, 3]}])
    parts, omissions = extract('test.xlsx', workbook(), options)
    assert [p.location for p in parts] == ['S1!R1C1', 'S1!R1C3', 'S1!R2C1', 'S1!R2C3']
    assert [p.field for p in parts[2:]] == ['이름', '휴대폰']
    assert any('1, 3' in text for text in omissions)
    assert any('시트 2' in text for text in omissions)
    assert 'private-title' not in json.dumps(omissions)


@pytest.mark.parametrize('selection', [
    [{'sheet': 4}], [{'sheet': 3}], [{'sheet': 1, 'columns': [5]}],
])
def test_missing_hidden_or_out_of_range_selection_fails_closed(selection):
    with pytest.raises(ValueError):
        extract('test.xlsx', workbook(), GuardOptions(excel_sheets=selection))


@pytest.mark.parametrize('selection', [
    [{'sheet': 0}], [{'sheet': True}], [{'sheet': '1'}],
    [{'sheet': 1}, {'sheet': 1}], [{'sheet': 1, 'columns': [0]}],
    [{'sheet': 1, 'columns': [1001]}], [{'sheet': 1, 'columns': [True]}],
])
def test_invalid_selection_is_not_silently_coerced(selection):
    with pytest.raises(ValidationError):
        GuardOptions(excel_sheets=selection)


def test_default_fingerprint_compatibility_and_selection_order():
    legacy = {'level': 'medium', 'fields': [{'field': '담당자', 'kind': 'name'}]}
    expected = hashlib.sha256((RULE_VERSION+json.dumps(legacy, ensure_ascii=False, separators=(',', ':'))).encode()).hexdigest()
    assert fingerprint(GuardOptions(**legacy)) == expected
    assert fingerprint(GuardOptions(**legacy, excel_sheets=[])) == expected
    one = GuardOptions(excel_sheets=[{'sheet': 2}, {'sheet': 1, 'columns': [3, 1, 1]}])
    two = GuardOptions(excel_sheets=[{'sheet': 1, 'columns': [1, 3]}, {'sheet': 2}])
    assert fingerprint(one) == fingerprint(two)
    assert fingerprint(one) != fingerprint(GuardOptions())


def test_selection_is_xlsx_only_and_empty_columns_mean_whole_selected_sheet():
    options = GuardOptions(excel_sheets=[{'sheet': 2}])
    parts, _ = extract('test.xlsx', workbook(), options)
    assert [p.text for p in parts] == ['unselected-sheet-marker']
    parts, _ = extract('test.csv', b'name,value\nexample,42', options)
    assert len(parts) == 4


class Recorder:
    def __init__(self): self.inputs = []
    async def run(self, context):
        self.inputs.append(context)
        return AgentResult(adapter='mock', output='selection received')


def test_api_selection_guard_storage_dispatch_restart_and_stale(tmp_path):
    agent = Recorder()
    root = tmp_path/'workspaces'
    base = '/api/projects/one'
    options = {'level': 'medium', 'excel_sheets': [{'sheet': 1, 'columns': [1, 3, 4]}]}
    with TestClient(create_app(root, agent)) as client:
        client.post('/api/projects', json={'name': 'one'})
        assert client.put(base+'/guard', json=options).status_code == 200
        response = client.post(base+'/materials?filename=example.xlsx', content=workbook())
        assert response.status_code == 201
        identity = response.json()['id']
        preview = client.get(base+'/materials/'+identity).json()['content']
        assert '홍*동' in preview and '010-****-2324' in preview
        assert '[비밀값]' in preview
        for forbidden in ['dummy-pass', 'internal-only-header', 'excluded-column-marker', 'unselected-sheet-marker', 'hidden-cell-marker', 'private-sheet-title']:
            assert forbidden not in preview + response.text
        result = client.post(base+'/commands', json={'text': 'test', 'material_ids': [identity]})
        assert result.status_code == 200
        assert agent.inputs[-1].files[0].content == preview
        run_id = result.json()['run_id']
    with TestClient(create_app(root, agent)) as client:
        assert client.get(base+'/guard').json()['excel_sheets'] == options['excel_sheets']
        assert client.get(base+'/materials/'+identity).json()['content'] == preview
        assert client.get(base+'/runs/'+run_id).json()['metadata']['guard']['state'] == 'dispatch_attempted'
        client.put(base+'/guard', json={'level': 'medium', 'excel_sheets': [{'sheet': 2}]})
        assert client.get(base+'/materials').json()['materials'][0]['stale']
        assert client.post(base+'/commands', json={'text': 'test', 'material_ids': [identity]}).status_code == 400
        count = len(client.get(base+'/materials').json()['materials'])
        client.put(base+'/guard', json={'excel_sheets': [{'sheet': 3}]})
        assert client.post(base+'/materials?filename=test.xlsx', content=workbook()).status_code == 400
        assert len(client.get(base+'/materials').json()['materials']) == count
