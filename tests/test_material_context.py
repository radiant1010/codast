import io
import json

import pytest

from app.main import create_app
from app.core.data_guard import GuardOptions
from app.models.schemas import Command, RulebookSettings


def setup_store(tmp_path):
    harness = create_app(tmp_path / 'ws').state.harness
    harness.projects.create('one')
    harness.save_rulebook('one', RulebookSettings(enabled=True, include_project_rules=True))
    return harness, harness.materials


def long_document():
    before = '\n\n'.join(f'일반 안내 {i}: ' + '배송 일정과 포장 안내입니다. ' * 30 for i in range(50))
    return before + '\n\n# 접근 정책\n\n권한은 관리자 승인 후 부여합니다.\n\n예외: 임시 계정은 하루 후 만료됩니다.\n\n# 부록\n\n' + before


def test_long_korean_lookup_keeps_heading_and_neighbor_after_restart(tmp_path):
    harness, store = setup_store(tmp_path)
    report = store.import_file('one', 'guide.md', long_document().encode())
    store = create_app(tmp_path / 'ws').state.harness.materials
    selected, audit = store.selected('one', [report['id']], query='권한은 어떻게 받아?')
    sent = selected[0].content
    assert all(text in sent for text in ['접근 정책', '관리자 승인', '임시 계정', '발췌'])
    assert len(sent) < len(long_document()) // 4
    delivery = audit['materials'][0]['delivery']
    assert delivery['mode'] == 'excerpt'
    assert delivery['sent_chars'] == len(sent)
    full = store.read('one', report['id'])['content']
    assert all(full[start:end] in sent for start, end in delivery['ranges'])


@pytest.mark.parametrize('heading', [True, False])
def test_prose_lookup_keeps_the_complete_section_or_falls_back(tmp_path, heading):
    _, store = setup_store(tmp_path)
    filler = 'General background.\n\n' * 500
    policy = ('Aurora access is restricted unless:\n\n'
              'the account has an approved sponsor; or\n\n'
              '| Condition | Limit |\n| --- | --- |\n| Temporary | One day |\n\n'
              'the account belongs to the emergency team.\n\n')
    text = filler + ('# Access\n\n' if heading else '') + policy + ('# Appendix\n\n' if heading else '') + filler
    report = store.import_file('one', 'guide.md', text.encode())
    selected, audit = store.selected('one', [report['id']], query='What is Aurora?')
    assert 'the account belongs to the emergency team.' in selected[0].content
    assert '| Temporary | One day |' in selected[0].content
    assert audit['materials'][0]['delivery']['mode'] == ('excerpt' if heading else 'full')


def test_prose_section_over_budget_is_not_partially_delivered(tmp_path):
    _, store = setup_store(tmp_path)
    text = '# Access\n\nAurora needs approval.\n\n' + 'More conditions.\n\n' * 500
    text += 'Exception: emergency accounts have no access.\n\n# Appendix\n\nEnd.'
    report = store.import_file('one', 'guide.md', text.encode())
    selected, audit = store.selected('one', [report['id']], query='What is Aurora?')
    assert audit['materials'][0]['delivery']['mode'] == 'full'
    assert selected[0].content == store.read('one', report['id'])['content']


def test_fifth_matching_section_is_kept_when_all_matches_fit(tmp_path):
    _, store = setup_store(tmp_path)
    text = ''.join(
        f'# Rule {i}\n\nAurora ' +
        ('permits standard access.' if i < 5 else 'prohibits temporary access without an approved sponsor.') + '\n\n'
        for i in range(1, 6)
    ) + '# Appendix\n\n' + 'Background only. ' * 600
    report = store.import_file('one', 'guide.md', text.encode())
    selected, audit = store.selected('one', [report['id']], query='What is Aurora?')
    assert audit['materials'][0]['delivery']['mode'] == 'excerpt'
    assert all(f'# Rule {i}' in selected[0].content for i in range(1, 6))
    assert 'prohibits temporary access without an approved sponsor.' in selected[0].content
    assert 'Background only.' not in selected[0].content


def test_later_matching_table_range_over_budget_keeps_full_text(tmp_path):
    _, store = setup_store(tmp_path)
    text = '| Product | Constraint |\n| --- | --- |\n| Aurora | standard use allowed |\n'
    text += ''.join(f'| other{i:03d} | ' + 'ordinary background ' * 8 + '|\n' for i in range(30))
    text += '| Aurora | temporary accounts are prohibited |\n'
    text += ''.join(f'| other{i:03d} | ' + 'ordinary background ' * 8 + '|\n' for i in range(30, 80))
    report = store.import_file('one', 'table.md', text.encode())
    selected, audit = store.selected('one', [report['id']], query='What is Aurora?', budget=3500)
    assert audit['materials'][0]['delivery']['mode'] == 'full'
    assert selected[0].content == store.read('one', report['id'])['content']
    assert 'temporary accounts are prohibited' in selected[0].content


def test_docx_table_label_does_not_end_a_prose_section(tmp_path):
    from docx import Document
    _, store = setup_store(tmp_path)
    doc = Document()
    for _ in range(30):
        doc.add_paragraph('General background. ' * 20)
    doc.add_heading('Access', level=2)
    doc.add_paragraph('Aurora access needs approval.')
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = 'Condition'
    table.cell(1, 0).text = 'Temporary access lasts one day.'
    doc.add_paragraph('Exception: emergency accounts have no access.')
    doc.add_heading('Appendix', level=2)
    doc.add_paragraph('Other information.')
    stream = io.BytesIO()
    doc.save(stream)
    report = store.import_file('one', 'guide.docx', stream.getvalue())
    for query in ('What is Aurora?', 'What is emergency?'):
        selected, audit = store.selected('one', [report['id']], query=query)
        assert audit['materials'][0]['delivery']['mode'] == 'excerpt'
        assert 'Aurora access needs approval.' in selected[0].content
        assert 'Temporary access lasts one day.' in selected[0].content
        assert 'Exception: emergency accounts have no access.' in selected[0].content


@pytest.mark.parametrize('query', ['', '전체 자료 요약해 줘', '권한 기능을 구현해', '권한 정책 검토', '이 자료로 개발해', '없는단어는 어디 있어?', '권한 몇?', 'List all permissions', '권한 요구 사항을 정리해 줘'])
def test_whole_document_requests_and_unmatched_search_keep_full_text(tmp_path, query):
    _, store = setup_store(tmp_path)
    report = store.import_file('one', 'guide.md', long_document().encode())
    selected, audit = store.selected('one', [report['id']], query=query)
    assert selected[0].content == store.read('one', report['id'])['content']
    assert audit['materials'][0]['delivery']['mode'] == 'full'


def test_large_table_lookup_retains_headers_and_guarded_values_only(tmp_path):
    _, store = setup_store(tmp_path)
    rows = ['항목,설명,이름'] + [f'item{i},일반 배송 안내,홍길동' for i in range(900)]
    rows[450] = '권한,관리자 승인,홍길동'
    report = store.import_file('one', 'table.csv', '\n'.join(rows).encode())
    selected, _ = store.selected('one', [report['id']], query='권한 설명 알려줘')
    assert '| 항목 | 설명 | 이름 |' in selected[0].content
    assert '| 권한 | 관리자 승인 | [이름] |' in selected[0].content
    assert '홍길동' not in selected[0].content
    with store.storage.connect() as db:
        indexed = '\n'.join(row[0] for row in db.execute('SELECT content FROM material_chunks'))
    assert '홍길동' not in indexed
    assert len(selected[0].content) < 6000


def test_index_scope_stale_policy_and_delete(tmp_path):
    harness, store = setup_store(tmp_path)
    harness.projects.create('two')
    own = store.import_file('one', 'guide.md', long_document().encode())
    other = store.import_file('two', 'guide.md', long_document().encode())
    with pytest.raises(FileNotFoundError):
        store.selected('one', [other['id']], query='권한은?')
    store.configure('one', GuardOptions(level='allow'))
    with pytest.raises(ValueError, match='설정'):
        store.selected('one', [own['id']], query='권한은?')
    store.delete('two', own['id'])
    store.delete('one', own['id'])
    with store.storage.connect() as db:
        assert db.execute('SELECT count(*) FROM material_chunks WHERE material=?', (own['id'],)).fetchone()[0] == 0
        assert db.execute('SELECT count(*) FROM material_chunks WHERE material=?', (other['id'],)).fetchone()[0] > 0


@pytest.mark.parametrize('client', ['mock', 'codex', 'claude'])
def test_context_budget_retrieval_and_common_rules(tmp_path, client):
    harness, store = setup_store(tmp_path)
    report = store.import_file('one', 'guide.md', long_document().encode())
    prepared = harness.prepare('one', Command(text='권한은 어떻게 받아?', material_ids=[report['id']], client=client))
    context = prepared[2]
    assert len(context.model_dump_json()) < harness.context.max_chars
    assert '임시 계정' in context.files[0].content
    assert any('구조 변경은 보수적으로' in rule.content for rule in context.rules)
    audit = harness.storage.run('one', prepared[0])['metadata']['guard']
    assert audit['materials'][0]['delivery']['mode'] == 'excerpt'


def test_docx_body_order_and_pptx_nested_groups(tmp_path):
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches
    _, store = setup_store(tmp_path)
    doc = Document()
    doc.add_heading('첫 번째 정책', level=1)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = '중간 표'
    doc.add_paragraph('마지막 문단')
    stream = io.BytesIO()
    doc.save(stream)
    report = store.import_file('one', 'order.docx', stream.getvalue())
    text = store.read('one', report['id'])['content']
    assert text.index('첫 번째 정책') < text.index('중간 표') < text.index('마지막 문단')
    assert '# 첫 번째 정책' in text
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape().shapes.add_group_shape()
    group.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text = '그룹 안의 정책'
    stream = io.BytesIO()
    deck.save(stream)
    report = store.import_file('one', 'group.pptx', stream.getvalue())
    assert '그룹 안의 정책' in store.read('one', report['id'])['content']


def test_v10_upload_is_reindexed_without_reextraction(tmp_path, monkeypatch):
    _, store = setup_store(tmp_path)
    report = store.import_file('one', 'guide.md', long_document().encode())
    with store.storage.connect() as db:
        db.execute('DROP TABLE material_chunks')
        db.execute('PRAGMA user_version=10')
    def forbidden(*args, **kwargs):
        raise AssertionError('Original extraction must not run')
    monkeypatch.setattr('app.core.materials.extract', forbidden)
    store = create_app(tmp_path / 'ws').state.harness.materials
    selected, audit = store.selected('one', [report['id']], query='권한은 어떻게 받아?')
    assert '관리자 승인' in selected[0].content
    assert audit['materials'][0]['delivery']['mode'] == 'excerpt'


def test_short_text_code_fences_and_full_workflow(tmp_path):
    harness, store = setup_store(tmp_path)
    short = store.import_file('one', 'short.md', b'# Retry\n\n```python\nretry = 3\n```')
    selected, audit = store.selected('one', [short['id']], query='What is retry?')
    assert selected[0].content == store.read('one', short['id'])['content']
    assert audit['materials'][0]['delivery']['mode'] == 'full'
    long = store.import_file('one', 'long.md', (long_document() + '\n\n```python\nretry = 3\n```').encode())
    selected, _ = store.selected('one', [long['id']], query='What is retry?')
    assert '```python\nretry = 3\n```' in selected[0].content
    with pytest.raises(ValueError, match='전체 전달'):
        harness.prepare('one', Command(text='권한은 어떻게 받아?', material_ids=[long['id']]), workflow_id='fixture')


@pytest.mark.parametrize('client', ['codex', 'claude'])
def test_reduced_context_reaches_both_cli_adapters(tmp_path, monkeypatch, client):
    import asyncio
    from app.models.schemas import AgentResult
    from app.llm.cli import CliAdapter
    harness, store = setup_store(tmp_path)
    report = store.import_file('one', 'long.md', long_document().encode())
    captured = []
    async def execute(self, context, *args):
        captured.append(json.loads(context))
        return AgentResult(adapter=client, output='fixture response')
    monkeypatch.setattr(CliAdapter, 'execute', execute)
    result = asyncio.run(harness.execute('one', Command(text='권한은 어떻게 받아?', material_ids=[report['id']], client=client)))
    assert len(captured[0]['files'][0]['content']) < 6000
    assert '관리자 승인' in captured[0]['files'][0]['content']
    run = harness.storage.run('one', result.run_id)
    assert run['metadata']['guard']['state'] == 'dispatch_attempted'


def test_multiple_attachments_share_remaining_context_budget(tmp_path):
    harness, store = setup_store(tmp_path)
    ids = [store.import_file('one', 'long.md', long_document().encode())['id'] for _ in range(6)]
    prepared = harness.prepare('one', Command(text='권한은 어떻게 받아?', material_ids=ids))
    assert len(prepared[2].model_dump_json()) <= harness.context.max_chars
    assert len(prepared[2].files) == 6
    assert all('관리자 승인' in part.content for part in prepared[2].files)


@pytest.mark.parametrize('tail,query,answer', [
    ('# Authorization\n\nAsk an administrator for approval.', 'What is authorization?', 'administrator'),
    ('| price | target |\n| --- | --- |\n| 500 | basic |\n| 900 | extended |', 'What is price?', '| 900 | extended |'),
])
def test_heading_and_header_hits_include_their_content(tmp_path, tail, query, answer):
    _, store = setup_store(tmp_path)
    report = store.import_file('one', 'guide.md', (long_document() + '\n\n' + tail).encode())
    selected, _ = store.selected('one', [report['id']], query=query)
    assert answer in selected[0].content


def test_all_query_terms_must_remain_in_the_excerpt(tmp_path):
    _, store = setup_store(tmp_path)
    words = ['alpha', 'beta', 'gamma', 'delta', 'epsilon']
    text = '\n\n'.join(word + '\n\n' + ('unrelated background ' * 60 + '\n\n') * 3 for word in words)
    report = store.import_file('one', 'guide.md', text.encode())
    selected, _ = store.selected('one', [report['id']], query='What is ' + ' '.join(words) + '?')
    assert all(word in selected[0].content for word in words)


def test_markdown_table_without_outer_pipes_retains_header(tmp_path):
    _, store = setup_store(tmp_path)
    text = 'general filler\n\n' * 700 + 'SKU | Price\n--- | ---\n'
    text += '\n'.join(f'widget{i} | {i+100}' for i in range(600))
    report = store.import_file('one', 'table.md', text.encode())
    selected, _ = store.selected('one', [report['id']], query='What is widget350?')
    assert 'widget350 | 450' in selected[0].content
    assert 'SKU | Price\n--- | ---' in selected[0].content


def test_deleted_material_is_not_resurrected_by_lazy_indexing(tmp_path, monkeypatch):
    _, store = setup_store(tmp_path)
    report = store.import_file('one', 'guide.md', long_document().encode())
    read = store.read
    def read_then_delete(project, identity):
        row = read(project, identity)
        store.delete(project, identity)
        return row
    monkeypatch.setattr(store, 'read', read_then_delete)
    with pytest.raises(FileNotFoundError):
        store.selected('one', [report['id']], query='권한은 어떻게 받아?')
    with store.storage.connect() as db:
        assert db.execute('SELECT count(*) FROM material_chunks').fetchone()[0] == 0


def test_soft_wrapped_paragraph_keeps_final_exception(tmp_path):
    _, store = setup_store(tmp_path)
    text = 'general filler\n\n' * 700 + 'Authorization policy:\n'
    text += ('Ordinary users may request approval. ' * 18 + '\n') * 10
    text += 'Exception: temporary users must never be granted access.'
    report = store.import_file('one', 'guide.md', text.encode())
    selected, _ = store.selected('one', [report['id']], query='What is authorization?')
    assert 'Exception: temporary users must never be granted access.' in selected[0].content


def test_reindex_is_safe_when_another_reader_already_indexed(tmp_path):
    from app.core.material_context import index_material
    _, store = setup_store(tmp_path)
    report = store.import_file('one', 'guide.md', long_document().encode())
    text = store.read('one', report['id'])['content']
    with store.storage.connect() as db:
        before = db.execute('SELECT count(*) FROM material_chunks').fetchone()[0]
        index_material(db, report['id'], text)
        assert db.execute('SELECT count(*) FROM material_chunks').fetchone()[0] == before


@pytest.mark.parametrize('suffix,prefix', [('.pdf','P'),('.pptx','S')])
def test_native_page_keeps_separated_values_notes_and_caption(tmp_path, monkeypatch, suffix, prefix):
    from app.core.data_guard import Part
    _, store = setup_store(tmp_path)
    kind = 'L' if suffix == '.pdf' else 'T'
    parts = [Part(f'{prefix}1{kind}{i}', 'Unrelated introduction ' * 30) for i in range(1,20)]
    parts.extend([Part(f'{prefix}2{kind}1','Aurora specification'),
                  Part(f'{prefix}2{kind}2','Valid from 2030-01-01'),
                  Part(f'{prefix}2T10R1C1','Variant'),Part(f'{prefix}2T10R1C2','Limit'),
                  Part(f'{prefix}2T10R2C1','Standard'),Part(f'{prefix}2T10R2C2','17'),
                  Part(f'{prefix}2{kind}11','Additional explanation'),
                  Part(f'{prefix}2{kind}12','Exception: temporary access expires after one hour.')])
    parts.extend(Part(f'{prefix}3{kind}{i}', 'Unrelated appendix ' * 30) for i in range(1,20))
    monkeypatch.setattr('app.core.materials.extract', lambda *args: (parts, []))
    report = store.import_file('one', 'native'+suffix, b'fixture')
    selected, audit = store.selected('one', [report['id']], query='What is Aurora?')
    assert audit['materials'][0]['delivery']['mode'] == 'excerpt'
    assert 'Valid from 2030-01-01' in selected[0].content
    assert '| Standard | 17 |' in selected[0].content
    assert 'Exception: temporary access expires after one hour.' in selected[0].content
    assert 'Unrelated appendix' not in selected[0].content


def test_old_native_upload_without_boundaries_uses_full_copy(tmp_path):
    _, store = setup_store(tmp_path)
    report = store.import_file('one', 'legacy.md', long_document().encode())
    with store.storage.connect() as db:
        old = json.loads(db.execute('SELECT report FROM materials WHERE id=?',(report['id'],)).fetchone()[0])
        old['label'] = old['label'].replace('.md','.pdf')
        db.execute('UPDATE materials SET report=? WHERE id=?',(json.dumps(old),report['id']))
    selected, audit = store.selected('one',[report['id']],query='권한은 어떻게 받아?')
    assert audit['materials'][0]['delivery']['mode'] == 'full'
    assert selected[0].content == store.read('one',report['id'])['content']


@pytest.mark.parametrize('suffix,prefix,marker', [
    ('.pdf', 'P', '# PDF 페이지 999'), ('.pptx', 'S', '# PPT 슬라이드 999'),
])
@pytest.mark.parametrize('spacing', ['', ' ', '\t', 'extra'])
def test_native_body_cannot_create_a_page_boundary(tmp_path, monkeypatch, suffix, prefix, marker, spacing):
    from app.core.data_guard import Part
    _, store = setup_store(tmp_path)
    kind = 'L' if suffix == '.pdf' else 'T'
    literal = marker.replace('# ', '#  ') if spacing == 'extra' else spacing+marker
    parts = [Part(f'{prefix}1{kind}1', 'Authorization policy'),
             Part(f'{prefix}1{kind}2', ('Literal example:\n' if not spacing else '')+literal),
             Part(f'{prefix}1{kind}3', 'Exception: temporary users are prohibited.'),
             Part(f'{prefix}2{kind}1', 'Unrelated appendix ' * 600)]
    monkeypatch.setattr('app.core.materials.extract', lambda *args: (parts, []))
    report = store.import_file('one', 'native'+suffix, b'fixture')
    selected, audit = store.selected('one', [report['id']], query='What is authorization?')
    assert audit['materials'][0]['delivery']['mode'] == 'excerpt'
    assert marker in selected[0].content
    assert 'Exception: temporary users are prohibited.' in selected[0].content
    assert 'Unrelated appendix' not in selected[0].content


@pytest.mark.parametrize('native', [False, True])
@pytest.mark.parametrize('reference', ['Aurora support follows EDU-202.', 'Aurora 지원은 EDU-202의 요구사항과 같이 제공한다.',
                                      'Aurora 지원은 EDU-<br>202의 요구사항과 같이 제공한다.'])
def test_requirement_references_follow_a_chain_and_stop_at_a_cycle(tmp_path, native, reference):
    from app.core.material_context import retrieve
    _, store = setup_store(tmp_path)
    titles = ['# PDF 페이지 ' + str(i) if native else '# Policy ' + str(i) for i in range(1, 5)]
    if not native:
        titles[-1] = '# Background'
    text = (titles[0] + '\n\n| 고유번호 | OPS-101 |\n| --- | --- |\n'
            '| 내용 | ' + reference + ' |\n\n' + titles[1] +
            '\n\n| 고유번호 | EDU-202 |\n| --- | --- |\n'
            '| 내용 | Training lasts at least 47 minutes; see EXC-303. |\n\n' + titles[2] +
            '\n\n| 고유번호 | EXC-303 |\n| --- | --- |\n'
            '| 내용 | Emergency accounts are excluded; see OPS-101. |\n\n' + titles[3] +
            '\n\n' + 'Unrelated background. ' * 500)
    report = store.import_file('one', 'chain.md', text.encode())
    full = store.read('one', report['id'])['content']
    with store.storage.connect() as db:
        sent, ranges = retrieve(db, report['id'], full, ['ops-101'], 6000, whole_pages=native)
    assert ranges == []
    assert sent == full
    assert 'Training lasts at least 47 minutes' in sent
    assert 'Emergency accounts are excluded' in sent


@pytest.mark.parametrize('reference', [
    'See MISSING-909.', '앞 절의 예외 조건을 따른다.', '앞<br>절의 예외 조건을 따른다.', 'See the previous section.',
    'See EDU-202 and Appendix A.', 'REQ-10 through REQ-12 apply.',
])
def test_unresolved_or_ambiguous_reference_delivers_full_text(tmp_path, reference):
    _, store = setup_store(tmp_path)
    text = '# OPS-101\n\nAurora needs approval. ' + reference
    text += '\n\n# Appendix\n\n' + 'Background. ' * 900
    report = store.import_file('one', 'uncertain.md', text.encode())
    selected, audit = store.selected('one', [report['id']], query='What is OPS-101?')
    assert audit['materials'][0]['delivery']['mode'] == 'full'
    assert selected[0].content == store.read('one', report['id'])['content']


def test_duplicate_requirement_definition_delivers_full_text(tmp_path):
    _, store = setup_store(tmp_path)
    text = '# OPS-101\n\nAurora follows EDU-202.\n\n'
    for name, condition in [('Training', '47 minutes'), ('Alternative', '93 minutes')]:
        text += f'# {name}\n\n| 고유번호 | EDU-202 |\n| --- | --- |\n| 내용 | {condition} |\n\n'
    text += '# Appendix\n\n' + 'Background. ' * 900
    report = store.import_file('one', 'ambiguous.md', text.encode())
    selected, audit = store.selected('one', [report['id']], query='What is OPS-101?')
    assert audit['materials'][0]['delivery']['mode'] == 'full'
    assert selected[0].content == store.read('one', report['id'])['content']


def test_reference_expansion_over_budget_preserves_full_text_and_blocks_dispatch(tmp_path):
    harness, store = setup_store(tmp_path)
    text = '# OPS-101\n\nAurora follows EDU-202.\n\n# Training\n\n'
    text += '| 고유번호 | EDU-202 |\n| --- | --- |\n| 내용 | Training lasts 47 minutes. |\n\n'
    text += 'Other training conditions.\n\n' * 1800
    text += '# Appendix\n\nBackground.'
    report = store.import_file('one', 'oversize.md', text.encode())
    selected, audit = store.selected('one', [report['id']], query='What is OPS-101?')
    assert audit['materials'][0]['delivery']['mode'] == 'full'
    assert selected[0].content == store.read('one', report['id'])['content']
    with pytest.raises(ValueError) as error:
        harness.prepare('one', Command(text='What is OPS-101?', material_ids=[report['id']]))
    assert '32,000' in str(error.value)
    assert '전체 전달' in str(error.value)
    assert harness.storage.runs('one', 10, 0) == []


def test_reference_definition_with_blank_table_cells_and_unicode_codes(tmp_path):
    _, store = setup_store(tmp_path)
    text = '# OPS-101\n\nAurora follows 교육‐202.\n\n# Training\n\n'
    text += '| 요구사항 고유번호 |  | 교육 - 202 |  |\n| --- | --- | --- | --- |\n'
    text += '| 내용 | Training must last at least 47 minutes. | | |\n\n'
    text += '# Appendix\n\n' + 'Background. ' * 900
    report = store.import_file('one', 'unicode.md', text.encode())
    selected, audit = store.selected('one', [report['id']], query='What is OPS-101?')
    assert audit['materials'][0]['delivery']['mode'] == 'full'
    assert 'Training must last at least 47 minutes.' in selected[0].content
    assert selected[0].content == store.read('one', report['id'])['content']


def test_requirement_code_range_keeps_intermediate_requirements_in_full_text(tmp_path):
    _, store = setup_store(tmp_path)
    text = '# OPS-101\n\nAurora follows REQ-10 through REQ-12.\n\n'
    for code, condition in [('REQ-10', 'Standard approval.'), ('REQ-11', 'Emergency users are excluded.'), ('REQ-12', 'Valid for 47 minutes.')]:
        text += f'# {code}\n\n{condition}\n\n'
    text += '# Appendix\n\n' + 'Background. ' * 900
    report = store.import_file('one', 'range.md', text.encode())
    selected, audit = store.selected('one', [report['id']], query='What is OPS-101?')
    assert audit['materials'][0]['delivery']['mode'] == 'full'
    assert selected[0].content == store.read('one', report['id'])['content']
    assert 'Emergency users are excluded.' in selected[0].content


@pytest.mark.parametrize('native', [False, True])
def test_value_lookup_without_requirement_id_keeps_competing_conditions(tmp_path, native):
    from app.core.material_context import retrieve
    _, store = setup_store(tmp_path)
    text = ('# PDF 페이지 1\n\n# OPS-101\n\n장애 접수 후 8시간 이내 대응 체계를 마련한다.\n\n'
            '# PDF 페이지 2\n\n# OPS-202\n\n장애 접수 후 4시간 이내 처리한다.\n\n'
            '# PDF 페이지 3\n\n' + 'General background. ' * 600)
    report = store.import_file('one', 'competing.md', text.encode())
    full = store.read('one', report['id'])['content']
    with store.storage.connect() as db:
        sent, ranges = retrieve(db, report['id'], full, ['8시간'], 6000, whole_pages=native)
    assert ranges == []
    assert sent == full
    assert '4시간' in sent


def test_unlinked_general_condition_outside_requirement_closure_keeps_full_text(tmp_path):
    _, store = setup_store(tmp_path)
    text = '# OPS-101\n\nSupport follows EDU-202.\n\n# EDU-202\n\nTraining lasts 1 hour weekly.\n\n'
    text += '# General support\n\n교육은 최소 2개월 이상, 주 2회 실시해야 한다.\n\n'
    text += '# Background\n\n' + 'Background. ' * 900
    report = store.import_file('one', 'general.md', text.encode())
    selected, audit = store.selected('one', [report['id']], query='What is OPS-101?')
    assert audit['materials'][0]['delivery']['mode'] == 'full'
    assert selected[0].content == store.read('one', report['id'])['content']
