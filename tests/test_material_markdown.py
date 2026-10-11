import io
import socket

import pytest

from app.main import create_app


def store_at(root):
    harness = create_app(root).state.harness
    harness.projects.create('one')
    return harness.materials


def test_guarded_table_is_saved_as_markdown_and_reused_after_restart(tmp_path, monkeypatch):
    store = store_at(tmp_path/'ws')
    def forbidden(*args, **kwargs):
        raise AssertionError('Conversion must stay local')
    monkeypatch.setattr(socket.socket, 'connect', forbidden)
    result = store.import_file('one', 'private-name.csv',
                               '이름,요청,비밀번호\n홍길동,배송 요청,example-pass'.encode())
    text = store.read('one', result['id'])['content']
    assert '| 이름 | 요청 | 비밀번호 |' in text
    assert '| [이름] | 배송 요청 | [비밀값] |' in text
    assert '홍길동' not in text and 'example-pass' not in text
    assert '[R2C' not in text
    restored = create_app(tmp_path/'ws').state.harness.materials
    monkeypatch.setattr('app.core.materials.extract', forbidden)
    selected, _ = restored.selected('one', [result['id']])
    assert selected[0].content == text


def test_markdown_keeps_blank_cells_and_literal_html(tmp_path):
    store = store_at(tmp_path/'ws')
    result = store.import_file('one', 'table.csv',
                               '항목,빈칸,설명\nA,,"<script>visible text</script>"\nB,,끝'.encode())
    text = store.read('one', result['id'])['content']
    assert '| A |  | <script>visible text</script> |' in text
    assert '| B |  | 끝 |' in text


def test_excel_selection_is_preserved_in_markdown(tmp_path):
    from openpyxl import Workbook
    from app.core.data_guard import GuardOptions, ExcelSheetSelection
    book = Workbook()
    sheet = book.active
    sheet.append(['선택', '제외', '설명'])
    sheet.append(['A', 'excluded-cell', '배송 요청'])
    sheet.append(['B', None, '=1+1'])
    hidden = book.create_sheet('hidden-source')
    hidden.sheet_state = 'hidden'
    hidden.append(['excluded-hidden'])
    stream = io.BytesIO()
    book.save(stream)
    store = store_at(tmp_path/'ws')
    store.configure('one', GuardOptions(excel_sheets=[ExcelSheetSelection(sheet=1, columns=[1, 3])]))
    result = store.import_file('one', 'table.xlsx', stream.getvalue())
    text = store.read('one', result['id'])['content']
    assert '| 선택 | 설명 |' in text
    assert '| A | 배송 요청 |' in text
    assert '| B |  |' in text
    assert all(value not in text for value in ['excluded-cell', 'excluded-hidden', '=1+1'])


def test_sparse_excel_does_not_spend_item_limit_on_blank_cells(tmp_path):
    from openpyxl import Workbook
    book = Workbook()
    book.active.cell(1, 1, 'item')
    book.active.cell(200, 100, 'far-cell')
    stream = io.BytesIO()
    book.save(stream)
    store = store_at(tmp_path/'ws')
    result = store.import_file('one', 'sparse.xlsx', stream.getvalue())
    text = store.read('one', result['id'])['content']
    assert 'far-cell' in text
    assert len(text.splitlines()[-1].split('|')) == 102


def test_pdf_table_keeps_an_entirely_blank_column(tmp_path):
    from tests.test_pdf_materials import TABLE, pdf_bytes
    store = store_at(tmp_path/'ws')
    page = TABLE.replace(b'(name)', b'()').replace(b'(Sample Person)', b'()')
    result = store.import_file('one', 'blank-column.pdf', pdf_bytes([page]))
    text = store.read('one', result['id'])['content']
    assert '|  | request |' in text
    assert '|  | Delivery request |' in text


def test_markdown_source_keeps_code_and_line_breaks(tmp_path):
    store = store_at(tmp_path/'ws')
    source = '# 제목\n\n```python\nprint("hello")\n```'
    result = store.import_file('one', 'source.md', source.encode())
    assert store.read('one', result['id'])['content'] == source


@pytest.mark.parametrize('suffix', ['.docx', '.pptx'])
def test_office_body_and_guarded_table_are_both_in_markdown(tmp_path, suffix):
    stream = io.BytesIO()
    if suffix == '.docx':
        from docx import Document
        document = Document()
        document.add_paragraph('Consultation')
        table = document.add_table(rows=2, cols=2)
        for row, values in zip(table.rows, [['name', 'request'], ['Sample Person', 'Delivery request']]):
            for cell, value in zip(row.cells, values):
                cell.text = value
        document.save(stream)
    else:
        from pptx import Presentation
        from pptx.util import Inches
        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, Inches(5), Inches(1)).text = 'Consultation'
        table = slide.shapes.add_table(2, 2, 0, Inches(1), Inches(5), Inches(2)).table
        for row, values in enumerate([['name', 'request'], ['Sample Person', 'Delivery request']]):
            for column, value in enumerate(values):
                table.cell(row, column).text = value
        document.save(stream)
    store = store_at(tmp_path/'ws')
    result = store.import_file('one', 'source'+suffix, stream.getvalue())
    text = store.read('one', result['id'])['content']
    assert 'Consultation' in text
    assert '| name | request |' in text
    assert '| [이름] | Delivery request |' in text
    assert 'Sample Person' not in text


def test_table_delimiters_and_multiline_cell_stay_in_one_cell(tmp_path):
    store = store_at(tmp_path/'ws')
    result = store.import_file('one', 'table.csv', b'item,value\nA,"left|right\nsecond line"')
    text = store.read('one', result['id'])['content']
    assert '| A | left\\|right' in text
    assert 'right<br>second line |' in text


def test_blocked_source_is_not_converted(tmp_path, monkeypatch):
    store = store_at(tmp_path/'ws')
    def forbidden(*args, **kwargs):
        raise AssertionError('Blocked sources must not be converted')
    monkeypatch.setattr('app.core.materials.to_markdown', forbidden, raising=False)
    result = store.import_file('one', 'key.txt', b'-----BEGIN PRIVATE KEY-----\nfake')
    assert result['status'] == 'blocked'
    assert store.read('one', result['id'])['content'] == ''


def test_conversion_failure_does_not_persist_source(tmp_path, monkeypatch):
    store = store_at(tmp_path/'ws')
    def fail(*args, **kwargs):
        raise RuntimeError('private-source-in-parser-error')
    monkeypatch.setattr('app.core.materials.to_markdown', fail, raising=False)
    with pytest.raises(ValueError) as error:
        store.import_file('one', 'table.csv', b'name,request\nSample Person,Delivery')
    assert 'private-source' not in str(error.value)
    assert store.list('one') == []
